from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.services.intelligence import score_asset
from app.services.standards_ingestor import parse_controls
from app.services.reporting import render_report

def test_scoring_is_bounded_and_resolved_vulnerabilities_do_not_count():
    asset = {'criticality':'critical', 'exposure':'internet'}
    assert score_asset(asset, [{'cvss_score':10,'status':'open'}])['score'] == 100
    assert score_asset(asset, [{'cvss_score':10,'status':'resolved'}])['score'] == 0
    assert score_asset({'criticality':'low'}, [{'severity':'high','status':'open'}])['score'] == 19.2

def test_standards_validate_before_any_writes():
    good = b'id,name,framework,asset_ids\nA1,Access,Custom,a1;a2\n'
    assert parse_controls('controls.csv', good)[0].asset_ids == ['a1','a2']
    with pytest.raises(ValueError): parse_controls('controls.json', b'[{}]')
    with pytest.raises(ValueError): parse_controls('controls.csv', good + b'A1,Duplicate,Custom,\n')
    with pytest.raises(ValueError): parse_controls('data.exe', b'bad')

def test_excel_export_treats_formula_like_names_as_text(tmp_path):
    path = tmp_path/'report.xlsx'
    render_report(path, 'asset', 'xlsx', [{'name':'=HYPERLINK("https://example.com")','score':42}])
    wb = load_workbook(path)
    assert wb.active['A2'].data_type == 's'
    assert wb.active['B2'].value == 42

def test_pdf_is_real_document(tmp_path):
    path = tmp_path/'report.pdf'
    render_report(path,'asset','pdf',[{'name':'<script> & server', 'id':'a1'}])
    assert path.read_bytes().startswith(b'%PDF-')

def test_compliance_is_graph_backed(client, admin_headers, monkeypatch):
    monkeypatch.setattr('app.services.intelligence.neo4j_client.run', lambda *a,**kw:[{'framework':'Custom','total_controls':4,'controls_met':1}])
    assert client.get('/api/v1/compliance/coverage',headers=admin_headers).json()[0]['coverage_pct'] == 25

def test_reports_generate_list_and_download(client, admin_headers, monkeypatch, tmp_path):
    monkeypatch.setattr('app.core.config.settings.REPORT_DIRECTORY', str(tmp_path))
    monkeypatch.setattr('app.api.v1.reports.report_rows', lambda kind:[{'name':'Server','id':'a1'}])
    result = client.post('/api/v1/reports/generate?report_type=asset&format=xlsx',headers=admin_headers)
    assert result.status_code == 201
    rid = result.json()['id']
    assert client.get('/api/v1/reports',headers=admin_headers).json()[0]['id'] == rid
    download = client.get(f'/api/v1/reports/{rid}/download',headers=admin_headers)
    assert download.status_code == 200
    assert load_workbook(BytesIO(download.content)).active['A2'].value == 'Server'
    assert client.post('/api/v1/reports/generate?report_type=invalid',headers=admin_headers).status_code == 422
    assert client.get(f'/api/v1/reports/{rid}/download').status_code == 401

def test_chat_routes_grounded_evidence_and_persists_history(client,admin_headers,monkeypatch):
    prompts=[]
    class Provider:
        def generate(self,prompt,model=None): prompts.append(prompt);return 'Server [a1] is present.'
    monkeypatch.setattr('app.agents.chat_assistant.get_model_provider',lambda:Provider())
    monkeypatch.setattr('app.agents.asset_agent.neo4j_client.list_assets',lambda **kw:[{'id':'a1','name':'Server'}])
    response=client.post('/api/v1/chat',json={'message':'List assets'},headers=admin_headers)
    assert response.status_code == 200
    data=response.json()
    assert data['citations'][0]['id']=='a1'
    assert data['agent_trace']==['orchestrator','asset_intelligence_agent','chat_assistant']
    response=client.post('/api/v1/chat',json={'message':'More about this asset','conversation_id':data['conversation_id']},headers=admin_headers)
    assert response.status_code==200
    assert 'List assets' in prompts[-1]
    assert client.post('/api/v1/chat',json={'message':' '},headers=admin_headers).status_code==422
    assert client.post('/api/v1/chat',json={'message':'assets','model':'not-enabled'},headers=admin_headers).status_code==422

def test_workflow_transitions_and_notifications(client,admin_headers):
    result=client.post('/api/v1/workflows',json={'title':'Patch server'},headers=admin_headers)
    assert result.status_code==201
    tid=result.json()['id']
    url=f'/api/v1/workflows/{tid}'
    assert client.patch(url,json={'status':'closed','expected_status':'open'},headers=admin_headers).status_code==409
    assert client.patch(url,json={'status':'in_progress','expected_status':'open'},headers=admin_headers).status_code==200
    assert client.patch(url,json={'status':'resolved','expected_status':'open'},headers=admin_headers).status_code==409
    note=client.get('/api/v1/notifications',headers=admin_headers).json()[0]
    assert note['read'] is False
    assert client.patch(f"/api/v1/notifications/{note['id']}/read",headers=admin_headers).json()['read'] is True

def test_risk_limit_is_bounded(client,admin_headers):
    assert client.get('/api/v1/risk/top?limit=-1',headers=admin_headers).status_code==422
    assert client.get('/api/v1/risk/top?limit=1001',headers=admin_headers).status_code==422


@pytest.mark.parametrize('status', ['mitigated', 'false_positive', 'accepted'])
def test_non_active_statuses_excluded(status):
    assert score_asset({'criticality':'critical'}, [{'cvss_score':10,'status':status}])['score']==0

def test_chat_does_not_invoke_forbidden_specialist(monkeypatch):
    from app.agents.orchestrator import run_orchestrator
    def forbidden(*args, **kwargs): raise AssertionError('Forbidden data retrieval')
    monkeypatch.setattr('app.agents.risk_agent.risk_tools.list_top_risks', forbidden)
    result=run_orchestrator('Show top risks', 'ComplianceOfficer')
    assert 'does not permit' in result['final_reply']

def test_private_resources_are_isolated(client,admin_headers,monkeypatch,tmp_path):
    monkeypatch.setattr('app.core.config.settings.REPORT_DIRECTORY', str(tmp_path))
    monkeypatch.setattr('app.api.v1.reports.report_rows', lambda kind:[])
    report=client.post('/api/v1/reports/generate?report_type=asset',headers=admin_headers).json()
    task=client.post('/api/v1/workflows',json={'title':'Private'},headers=admin_headers).json()
    client.post('/api/v1/admin/users',headers=admin_headers,json={'email':'other@csos.com','password':'other-user-password','full_name':'Other','role':'Admin','is_active':True})
    token=client.post('/api/v1/auth/login',json={'email':'other@csos.com','password':'other-user-password'}).json()['access_token']
    other={'Authorization':'Bearer '+token}
    assert client.get('/api/v1/reports',headers=other).json()==[]
    assert client.get(f"/api/v1/reports/{report['id']}/download",headers=other).status_code==404
    assert client.patch(f"/api/v1/workflows/{task['id']}",headers=other,json={'status':'in_progress','expected_status':'open'}).status_code==404


def test_bootstrap_preserves_admin_role_edits(db_session):
    from app.models.user import Role
    from app.services.bootstrap import seed_identity_data
    role = db_session.query(Role).filter_by(name='Analyst').one()
    role.permissions = []
    db_session.commit()
    seed_identity_data(db_session)
    assert role.permissions == []

def test_structured_pdf_controls(tmp_path):
    from reportlab.pdfgen import canvas
    path=tmp_path/'standard.pdf'
    pdf=canvas.Canvas(str(path))
    pdf.drawString(50,750,'id | name | framework')
    pdf.drawString(50,725,'AC-1 | Access control | Custom')
    pdf.save()
    controls=parse_controls(path.name,path.read_bytes())
    assert controls[0].name=='Access control'
    assert controls[0].status=='not_implemented'

def test_standards_invalid_input_does_not_write_graph(client,admin_headers,monkeypatch):
    def unexpected(*args,**kwargs): raise AssertionError('Invalid input wrote graph')
    monkeypatch.setattr('app.api.v1.standards.neo4j_client.run',unexpected)
    response=client.post('/api/v1/standards/upload',headers=admin_headers,files={'file':('controls.json',b'[{}]','application/json')})
    assert response.status_code==422

def test_manual_standard_saves_and_records_history(client,admin_headers,monkeypatch):
    calls=[]
    monkeypatch.setattr('app.api.v1.standards.neo4j_client.run',lambda *args,**kwargs:calls.append((args,kwargs)))
    response=client.post('/api/v1/standards/controls',headers=admin_headers,json={'id':'AC-1','name':'Access review','framework':'Custom'})
    assert response.status_code==201
    assert response.json()['controls_parsed']==1
    assert len(calls)==1
    assert client.get('/api/v1/standards',headers=admin_headers).json()[0]['status']=='processed'


def test_chat_discards_history_from_previous_role(client,admin_headers,db_session,monkeypatch):
    import uuid
    from app.models.conversation import Conversation
    user=client.get('/api/v1/auth/me',headers=admin_headers).json()
    conversation=Conversation(user_id=uuid.UUID(user['id']),messages=[{'role':'assistant','content':'previous role evidence','access_role':'Executive'}])
    db_session.add(conversation);db_session.commit()
    seen=[]
    def run(query,role,history,model):
        seen.extend(history)
        return {'final_reply':'answer','agent_trace':[]}
    monkeypatch.setattr('app.api.v1.chat.run_orchestrator',run)
    response=client.post('/api/v1/chat',headers=admin_headers,json={'message':'assets','conversation_id':str(conversation.id)})
    assert response.status_code==200
    assert seen==[]
