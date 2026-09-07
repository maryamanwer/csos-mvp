from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
from app.graph.neo4j_client import neo4j_client
from app.services.intelligence import assessments, coverage

def report_rows(kind):
    if kind == 'risk': return assessments()
    if kind == 'compliance': return coverage()
    assets = []
    offset = 0
    while True:
        page = neo4j_client.list_assets(offset=offset, limit=1000)
        assets.extend(page)
        if len(page) < 1000: break
        offset += 1000
        if offset >= 50000: raise ValueError('Export exceeds 50,000 assets; narrow your inventory first')
    return [{k: a.get(k, '') for k in ('id', 'name', 'type', 'criticality', 'environment', 'owner')} for a in assets]

def render_report(path: Path, kind: str, fmt: str, rows: list[dict]):
    columns = list(rows[0]) if rows else ['result']
    data = rows or [{'result': 'No records available'}]
    if fmt == 'xlsx':
        wb = Workbook()
        ws = wb.active
        ws.title = kind.title()
        ws.append(columns)
        for row in data:
            ws.append([v if isinstance(v, (int, float)) else str(v or '') for v in (row.get(c) for c in columns)])
        # All imported text is literal, including strings beginning with =, +, -, or @.
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(70, max(16, max(len(str(c.value or '')) for c in col) + 2))
        wb.save(path)
    else:
        styles = getSampleStyleSheet()
        body = styles['BodyText']
        body.fontSize = 8
        body.leading = 11
        table_data = [[Paragraph(escape(str(c).replace('_', ' ').title()), body) for c in columns]]
        for row in data:
            table_data.append([Paragraph(escape(str(row.get(c, '') or '')), body) for c in columns])
        table = LongTable(table_data, repeatRows=1, splitInRow=1, colWidths=[770 / len(columns)] * len(columns))
        table.setStyle(TableStyle([('BACKGROUND', (0,0),(-1,0),colors.HexColor('#dbe7f2')),
                                   ('VALIGN',(0,0),(-1,-1),'TOP'), ('GRID',(0,0),(-1,-1),.25,colors.lightgrey),
                                   ('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=35, rightMargin=35).build([
            Paragraph('CSOS — ' + kind.title() + ' report', styles['Title']),
            Paragraph(f'{len(rows)} records. Risk and compliance exports are bounded snapshots (up to 5,000 assets / 1,000 frameworks).', body),
            Spacer(1, 14), table])
