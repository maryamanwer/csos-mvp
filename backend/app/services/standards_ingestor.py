import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile
from typing import Literal
from pydantic import BaseModel, Field, ConfigDict, field_validator
from app.core.config import settings

class ControlInput(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)
    id: str = Field(min_length=1, max_length=100)
    name: str = Field(min_length=1, max_length=255)
    framework: str = Field(min_length=1, max_length=255)
    description: str = Field(default='', max_length=8000)
    status: Literal['not_implemented', 'partial', 'implemented'] = 'not_implemented'
    asset_ids: list[str] = Field(default_factory=list, max_length=1000)
    @field_validator('asset_ids', mode='before')
    @classmethod
    def split_ids(cls, value):
        return [v.strip() for v in value.split(';') if v.strip()] if isinstance(value, str) else value

def parse_controls(filename, content):
    suffix = Path(filename).suffix.lower()
    if len(content) > 10 * 1024 * 1024: raise ValueError('Maximum upload size is 10 MB')
    if suffix in ('.xlsx', '.docx'):
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Expanded document exceeds 50 MB')
    if suffix == '.csv':
        rows = list(csv.DictReader(StringIO(content.decode('utf-8-sig'))))
    elif suffix == '.json':
        rows = json.loads(content)
        if isinstance(rows, dict): rows = rows.get('controls')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            values = wb.active.iter_rows(values_only=True)
            header = next(values, ())
            rows = []
            for values_row in values:
                if any(v is not None for v in values_row):
                    rows.append({str(k).strip(): v for k, v in zip(header, values_row) if k and v is not None})
                if len(rows) > settings.MAX_IMPORT_ROWS: raise ValueError('Too many controls')
        finally: wb.close()
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted: raise ValueError('Encrypted PDFs are not supported')
        if len(reader.pages) > 100: raise ValueError('PDF exceeds 100 pages')
        lines = []
        for page in reader.pages:
            lines.extend(line.strip() for line in (page.extract_text() or '').splitlines() if '|' in line)
        if not lines: raise ValueError('PDF must contain selectable text with pipe-separated control rows; scanned PDFs require OCR outside CSOS')
        header = [v.strip() for v in lines[0].split('|')]
        if not {'id', 'name', 'framework'}.issubset(header):
            raise ValueError('PDF header must include id | name | framework')
        rows = []
        for line in lines[1:]:
            values = [v.strip() for v in line.split('|')]
            if values == header: continue
            if len(values) != len(header): raise ValueError('PDF row does not match its header')
            rows.append(dict(zip(header, values)))
    elif suffix == '.docx':
        from docx import Document
        doc = Document(BytesIO(content))
        rows = []
        for table in doc.tables:
            header = [cell.text.strip() for cell in table.rows[0].cells]
            rows.extend(dict(zip(header, [c.text.strip() for c in row.cells])) for row in table.rows[1:])
    else:
        raise ValueError('Use CSV, XLSX, JSON, a structured text PDF, or a Word control table')
    if not isinstance(rows, list) or not rows: raise ValueError('No controls found')
    if len(rows) > settings.MAX_IMPORT_ROWS: raise ValueError('Too many controls')
    controls = [ControlInput.model_validate(row) for row in rows]
    keys = [(c.framework, c.id) for c in controls]
    if len(set(keys)) != len(keys): raise ValueError('Duplicate control IDs within a framework')
    return controls
